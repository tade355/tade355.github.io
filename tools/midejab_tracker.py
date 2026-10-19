"""Append collections and payments to the Midejab Excel tracker.

    python tools/midejab_tracker.py in.xlsx out.xlsx entries.json

entries.json is a list, oldest first, for example:
    [{"date": "2026-10-19", "staff": "Christopher", "ago": 955, "rate": 1900},
     {"date": "2026-10-19", "staff": "Christopher", "pms": 20, "rate": 1385},
     {"date": "2026-10-20", "staff": "Payment", "payment": 5000000}]
"ago" may be a formula string such as "=700+990". One litre type or one payment per row, as in the sheet.

The sheet is the user's own layout (columns C to L: Date, Staff/Company, AGO, PMS, Rates, costs,
Payments, Balance), with a formula-driven balance and placeholder rows to the right of the last
entry. New entries go in the first empty placeholder rows, the placeholder dates below them are
renumbered day by day, the TOTAL formulas are extended to the last entry and the new rows are
unhidden. Check the result in the rendered picture before sending it.
"""
import datetime, json, re, sys
import openpyxl

src, dst, entries_file = sys.argv[1:4]
wb = openpyxl.load_workbook(src)
ws = wb.active
entries = json.load(open(entries_file))

FIRST, TOTAL_LABEL = 5, 'TOTAL'
total_row = next(r for r in range(FIRST, ws.max_row + 1) if ws.cell(r, 3).value == TOTAL_LABEL)
last_row = max(r for r in range(FIRST, total_row) if any(ws.cell(r, c).value not in (None, '') for c in (4, 5, 6, 11)))
placeholder_end = total_row - 2  # last dated placeholder row before the totals block

row = last_row
for e in entries:
    row += 1
    if row > placeholder_end:
        raise SystemExit('no placeholder rows left in the sheet')
    ws.cell(row, 3).value = datetime.datetime.fromisoformat(e['date'])
    ws.cell(row, 4).value = e.get('staff')
    ws.cell(row, 5).value = e.get('ago')
    ws.cell(row, 6).value = e.get('pms')
    ws.cell(row, 7).value = e.get('rate')
    ws.cell(row, 11).value = e.get('payment')
    ws.row_dimensions[row].hidden = False

# keep one placeholder per calendar day after the last entry
day = ws.cell(row, 3).value + datetime.timedelta(days=1)
for r in range(row + 1, placeholder_end + 1):
    ws.cell(r, 3).value = day
    day += datetime.timedelta(days=1)

for c in range(5, 12):  # extend the end of every SUM range in the TOTAL row
    v = ws.cell(total_row, c).value
    if isinstance(v, str) and v.startswith('=SUM('):
        ws.cell(total_row, c).value = re.sub(r':([A-Z]+)\d+\)', lambda m: f':{m.group(1)}{row})', v)

ws['C2'] = re.sub(r'– [^)]+\)', '– ' + ws.cell(row, 3).value.strftime('%-d %b %Y') + ')', ws['C2'].value)
wb.save(dst)
print(f'added {len(entries)} rows, last entry row {row}; totals row {total_row}')
